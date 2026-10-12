from pathlib import Path
from datetime import datetime,date,time
import re
from openpyxl import load_workbook
from docx import Document

ROOT=Path(__file__).resolve().parent
MESES=['enero','febrero','marzo','abril','mayo','junio','julio','agosto','septiembre','octubre','noviembre','diciembre']

def text(v):
    return '' if v is None else str(v)

def required(row,key):
    v=row.get(key)
    if v is None or not str(v).strip(): raise ValueError(f'Falta {key}.')
    return v

def fecha(v):
    if not isinstance(v,(date,datetime)): raise ValueError('Introduce fechas de Excel reales, no texto.')
    return v.strftime('%d/%m/%Y')

def cell(c,value):
    # Retain cell properties and paragraph/run styles; do not recreate tables.
    ps=c.paragraphs
    for p in ps:
        for r in p.runs:r.text=''
    p=ps[0]
    if p.runs:p.runs[0].text=text(value)
    else:p.add_run(text(value))

def cells(row):
    out=[];seen=set()
    for c in row.cells:
        if c._tc not in seen:out.append(c);seen.add(c._tc)
    return out

def fill_rows(rows,values):
    if len(values)>len(rows):raise ValueError(f'La plantilla admite {len(rows)} filas; se solicitaron {len(values)}. No se amplía para proteger el diseño.')
    for i,r in enumerate(rows):
        for j,c in enumerate(cells(r)):cell(c,values[i][j] if i<len(values) and j<len(values[i]) else '')

def replace(doc,old,new):
    # Replacement across split runs while retaining unaffected run formatting.
    roots=[doc.element]
    for sec in doc.sections:
        roots += [sec.header._element,sec.footer._element,sec.first_page_header._element,sec.first_page_footer._element,sec.even_page_header._element,sec.even_page_footer._element]
    from docx.oxml.ns import qn
    for root in roots:
        for p in root.iter(qn('w:p')):
            ts=list(p.iter(qn('w:t')))
            combined=''.join(t.text or '' for t in ts)
            while old in combined:
                start=combined.index(old);end=start+len(old);pos=0;written=False
                for t in ts:
                    value=t.text or '';a=pos;b=pos+len(value);pos=b
                    if b<=start or a>=end:continue
                    t.text=value[:max(0,start-a)]+(text(new) if not written else '')+value[max(0,end-a):]
                    t.set(qn('xml:space'),'preserve');written=True
                combined=''.join(t.text or '' for t in ts)
                if old in text(new):break

class Base:
    def __init__(self,path):
        w=load_workbook(path,data_only=True)
        self.data={}
        for s in w:
            rows=list(s.values);headers=rows[0]
            self.data[s.title]=[dict(zip(headers,r)) for r in rows[1:] if any(v is not None for v in r)]
    def rows(self,s,key=None,value=None):
        if s not in self.data:raise ValueError(f'Falta la hoja {s}.')
        return [r for r in self.data[s] if key is None or r.get(key)==value]
    def one(self,s,key,value):
        rs=self.rows(s,key,value)
        if len(rs)!=1:raise ValueError(f'{s}: {key}={value} debe identificar exactamente una fila.')
        return rs[0]
    def person(self,id):return self.one('Personal','persona_id',id)
    def name(self,id):return required(self.person(id),'nombre_completo')

    def render(self,job):
        pid=required(job,'plantilla_id');p=self.one('Plantillas','plantilla_id',pid)
        filename=required(p,'archivo_word')
        if Path(filename).name!=filename:raise ValueError('Nombre de plantilla no válido.')
        doc=Document(ROOT/'plantillas'/filename)
        e=self.one('Eventos','evento_id',required(job,'evento_id'))
        if e.get('estado_revision')!='Revisado':raise ValueError('El evento debe estar Revisado antes de generarlo.')
        st=self.one('Estaciones','estacion_id',required(e,'estacion_id'))
        station=required(st,'numero_estacion');day=required(e,'fecha');fecha(day)
        replace(doc,'E-02585',station)
        people=[]
        for a in sorted(self.rows('Asistentes','evento_id',e['evento_id']),key=lambda r:r.get('orden') or 0):
            if a.get('asistencia')=='Confirmada':
                person=self.person(a['persona_id'])
                if person.get('estacion_id')!=st['estacion_id']:raise ValueError('Asistente de otra estación.')
                people.append(person)
        if pid in ('P04','P05','P06','P07'):
            hour=required(e,'hora')
            if not isinstance(hour,(time,datetime)):raise ValueError('La hora debe ser una hora real de Excel.')
            kind=required(e,'tipo_reunion')
            if kind not in ('Ordinaria','Extraordinaria'):raise ValueError('Tipo de reunión no válido.')
            for t in doc.tables:
                if not t.rows:continue
                first=cells(t.rows[0])[0].text
                if first=='Tipo de Reunión:':
                    cell(cells(t.rows[0])[1],f'Ordinaria: {"X" if kind=="Ordinaria" else "____"}    Extraordinaria: {"X" if kind=="Extraordinaria" else "____"}')
                    cs=cells(t.rows[1]);cell(cs[1],fecha(day));cell(cs[3],hour.strftime('%H:%M'));cell(cs[-1],required(e,'lugar'))
                    cell(cells(t.rows[2])[1],required(e,'objetivo'))
                    if len(t.rows)>3:
                        role=next(i for i,r in enumerate(t.rows) if cells(r)[0].text.startswith('Convocante:'))
                        ag=sorted(self.rows('Agenda','evento_id',e['evento_id']),key=lambda x:x.get('orden') or 0)
                        if not ag:raise ValueError('Falta la agenda del evento.')
                        for i,r in enumerate(t.rows[5:role]):
                            cs=cells(r)
                            if i<len(ag):
                                a=ag[i];vals=[required(a,'tema'),self.name(required(a,'responsable_id')),f'{required(a,"duracion_minutos")} MIN.']
                            else:vals=['','','']
                            for c,v in zip(cs[-3:],vals):cell(c,v)
                        if len(ag)>role-5:raise ValueError('La agenda excede la capacidad de la plantilla.')
                        cell(cells(t.rows[role])[1],'\n'.join(self.name(required(e,k)) for k in ('convocante_id','facilitador_id','secretario_id')))
                        agreements=self.rows('Acuerdos','evento_id',e['evento_id'])
                        if any(a.get('confirmado')!='Sí' for a in agreements):raise ValueError('Hay acuerdos pendientes de confirmar.')
                        vals=[]
                        for a in sorted(agreements,key=lambda x:x.get('orden') or 0):
                            owner=self.name(a['responsable_id']) if a.get('responsable_id') else required(a,'responsable_texto')
                            vals.append([required(a,'acuerdo'),owner,fecha(a['fecha_entrega']) if a.get('fecha_entrega') else ''])
                        fill_rows(t.rows[role+2:],vals)
                elif first in ('NOMBRE','N°'):
                    if not people:raise ValueError('No hay asistentes confirmados.')
                    if first=='N°':vals=[[i,required(x,'nombre_completo'),required(x,'puesto'),text(x.get('telefono')),''] for i,x in enumerate(people,1)]
                    else:vals=[[required(x,'nombre_completo'),required(x,'puesto'),text(x.get('cargo_sasisopa')),text(x.get('correo')),text(x.get('telefono'))] for x in people]
                    fill_rows(t.rows[1:],vals)
            replace(doc,'13/ENE/2026',fecha(day))
        elif pid=='P02':
            s=self.one('Sesiones','evento_id',e['evento_id']);t=doc.tables[1]
            if not people:raise ValueError('No hay asistentes confirmados.')
            if len(people)>10:raise ValueError('El plan admite hasta 10 participantes.')
            cell(t.cell(0,1),fecha(day));cell(t.cell(0,2),f'Duración: {required(s,"duracion_minutos")} MIN.');cell(t.cell(0,3),'Lugar: '+required(e,'lugar'))
            for row,key,source in [(1,'tema',e),(2,'objetivo_general',s),(5,'contenido_tematico',s)]:cell(t.cell(row,1),required(source,key))
            cell(t.cell(3,1),f'Todo el personal de la estación {station}\n\n'+'\n\n'.join(f'{i}. {required(x,"nombre_completo")}' for i,x in enumerate(people,1))+'\n\n'+text(s.get('material_entregado')))
            cell(t.cell(4,1),required(s,'numero_sesiones'));cell(t.cell(4,3),len(people));cell(t.cell(6,1),self.name(required(s,'instructor_id')));cell(t.cell(8,1),text(s.get('evaluacion_diagnostica')))
        elif pid=='P03':
            n=self.one('Notas','evento_id',e['evento_id'])
            if n.get('confirmada')!='Sí':raise ValueError('La nota debe estar confirmada.')
            if n.get('hubo_cambios_personal') not in ('Sí','No') or n.get('hubo_cambios_estructura') not in ('Sí','No'):raise ValueError('Confirma si hubo cambios.')
            ps=doc.paragraphs
            body=next(x for x in ps if x.text.startswith('En la estación'))
            cell_like_paragraph(body,required(n,'texto_nota'))
            for p in ps:
                if p.text=='13 de Enero del 2026':cell_like_paragraph(p,f'{day.day} de {MESES[day.month-1]} de {day.year}')
                elif p.text=='ABEL E. TREVIÑO GALLEGOS':cell_like_paragraph(p,self.name(required(n,'firmante_id')))
                elif p.text=='Representante Técnico':cell_like_paragraph(p,required(self.person(n['firmante_id']),'puesto'))
        elif pid=='P01':
            a=self.one('Acuses','acuse_id',required(job,'acuse_id'))
            if a.get('evento_id')!=e['evento_id']:raise ValueError('El acuse pertenece a otro evento.')
            replace(doc,'13/ENE /2026',fecha(required(a,'fecha')));replace(doc,'Acuse número: _____','Acuse número: '+text(required(a,'numero_acuse')))
            replace(doc,'Documento: Tríptico.','Documento: '+text(required(a,'tipo_documento')))
            replace(doc,'No controlada.',text(required(a,'tipo_copia'))+'.')
            ds=sorted(self.rows('Documentos_Entregados','acuse_id',a['acuse_id']),key=lambda x:x.get('orden') or 0)
            if not ds:raise ValueError('Faltan documentos entregados.')
            fill_rows(doc.tables[0].rows[1:-1],[[i,text(x.get('codigo')),text(x.get('revision')),required(x,'nombre_documento'),required(x,'numero_hojas')] for i,x in enumerate(ds,1)])
            cell(doc.tables[0].rows[-1].cells[-1],sum(int(required(x,'numero_hojas')) for x in ds))
            rs=self.rows('Recepciones','acuse_id',a['acuse_id'])
            if not rs or any(x.get('recepcion_confirmada')!='Sí' for x in rs):raise ValueError('Faltan recepciones confirmadas.')
            vals=[]
            for x in rs:
                person=self.person(x['persona_id']);copies=int(required(x,'cantidad_copias'))
                if copies<=0:raise ValueError('Las copias deben ser positivas.')
                vals.append([copies,required(person,'puesto'),required(person,'nombre_completo'),'',fecha(required(x,'fecha_recepcion'))])
            fill_rows(doc.tables[1].rows[1:-1],vals);cell(doc.tables[1].rows[-1].cells[0],'Total: '+str(sum(x[0] for x in vals)))
            person=self.person(required(a,'entrego_id'));cell(doc.tables[2].cell(0,0),f'Entregó:\nNombre: {required(person,"nombre_completo")}    Puesto: {required(person,"puesto")}\nFirma: __________________________')
        else:raise ValueError('Plantilla no compatible.')
        # Remove old station-owner text from headers/body where it is not part of edited fields.
        replace(doc,'ABEL TREVINO GONZALEZ',required(st,'razon_social'))
        return doc

    def generate(self,out):
        jobs=[j for j in self.rows('Generar') if j.get('generar')=='Sí']
        if not jobs:raise ValueError('Marca Sí en la columna generar de la hoja Generar.')
        prepared=[];names=set()
        for j in jobs:
            name=text(required(j,'nombre_salida'))
            if re.search(r'[<>:"/\\|?*\x00-\x1f]',name) or name.endswith((' ','.')) or name.split('.')[0].upper() in {'CON','PRN','AUX','NUL',*[f'COM{i}' for i in range(1,10)],*[f'LPT{i}' for i in range(1,10)]}:raise ValueError('Nombre de salida no válido.')
            if name.casefold() in names:raise ValueError('Nombres de salida duplicados.')
            names.add(name.casefold());prepared.append((name,self.render(j)))
        out=Path(out);out.mkdir(parents=True,exist_ok=True)
        paths=[out/(name+'.docx') for name,_ in prepared]
        if any(p.exists() for p in paths):raise ValueError('Ya existe un documento de salida. Elige una carpeta nueva para no sobrescribirlo.')
        for p,(_,doc) in zip(paths,prepared):doc.save(p)
        return paths

def cell_like_paragraph(p,value):
    for r in p.runs:r.text=''
    if p.runs:p.runs[0].text=text(value)
    else:p.add_run(text(value))

def export_pdf(paths):
    import pythoncom
    import win32com.client
    pythoncom.CoInitialize()
    word=None
    try:
        word=win32com.client.DispatchEx('Word.Application')
        word.Visible=False
        for path in paths:
            pdf=path.with_suffix('.pdf')
            if pdf.exists():raise ValueError('El PDF ya existe: '+pdf.name)
            doc=word.Documents.Open(str(path.resolve()),ReadOnly=True,AddToRecentFiles=False)
            try:
                doc.ExportAsFixedFormat(str(pdf.resolve()),17)
                if not pdf.is_file() or pdf.stat().st_size==0:
                    raise RuntimeError('Word no creó el PDF: '+pdf.name)
            finally:doc.Close(False)
    finally:
        try:
            if word is not None:word.Quit()
        finally:pythoncom.CoUninitialize()
