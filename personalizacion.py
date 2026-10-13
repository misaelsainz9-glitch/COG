"""Encabezados y firmas en los espacios existentes de las plantillas."""
from pathlib import Path
import re
from PIL import Image as PILImage
from docx.image.image import Image
from docx.oxml.ns import qn
from docx.shared import Cm
from motor import cells,cell_like_paragraph,fecha,MESES


def validated_image(path,label):
    path=Path(path)
    if not path.is_file():raise ValueError(f'No se encuentra {label}: {path.name}')
    if path.suffix.lower() not in ('.png','.jpg','.jpeg'):
        raise ValueError(f'{label}: usa PNG o JPG.')
    try:
        with PILImage.open(path) as probe:probe.verify()
        with PILImage.open(path) as probe:probe.load()
        image=Image.from_file(str(path))
    except Exception as exc:raise ValueError(f'No se puede leer {label}; selecciona una imagen PNG o JPG válida.') from exc
    if image.px_width<=0 or image.px_height<=0:raise ValueError(f'{label}: dimensiones no válidas.')
    return path,image


def fit(image,width,height):
    ratio=min(width/image.px_width,height/image.px_height)
    return max(1,int(image.px_width*ratio)),max(1,int(image.px_height*ratio))


def apply_header(doc,address='',logo='',revision='',revision_date=None):
    revision=revision.strip()
    if revision and not re.fullmatch(r'[0-9]+',revision):
        raise ValueError('La revisión del encabezado debe ser un número entero, por ejemplo 0, 1 o 2.')
    if revision:revision=str(int(revision))
    logo_image=validated_image(logo,'el logo') if logo else None
    count=0;addresses=0
    for rel in doc.part.rels.values():
        if rel.is_external or not rel.reltype.endswith('/header'):continue
        part=rel.target_part;root=part.element
        paragraphs=list(root.iter(qn('w:p')))
        for para in paragraphs:
            ts=list(para.iter(qn('w:t')));value=''.join(t.text or '' for t in ts)
            if address and value.strip().startswith('CARRETERA MONTERREY-REYNOSA'):
                # Replace header address text only; preserve its paragraph, drawings and run properties.
                ts[0].text=address.strip();ts[0].set(qn('xml:space'),'preserve')
                for t in ts[1:]:t.text=''
                addresses+=1
        if revision or revision_date:
            flat=list(root.iter(qn('w:t')))
            if revision:
                for node in flat:
                    if node.text and re.search(r'\.R[0-9]+$',node.text):
                        node.text=re.sub(r'\.R[0-9]+$',f'.R{revision}',node.text)
            for i,node in enumerate(flat):
                if node.text=='REVISIÓN:' and i+1<len(flat):
                    if revision:flat[i+1].text=revision.strip()
                    if revision_date:
                        fecha(revision_date)
                        if i+3>=len(flat) or flat[i+2].text!='JULIO' or flat[i+3].text!='2018':
                            raise ValueError('La fecha de revisión del encabezado no tiene la estructura esperada.')
                        flat[i+2].text=MESES[revision_date.month-1].upper();flat[i+3].text=str(revision_date.year)
        if logo_image:
            path,image=logo_image
            for drawing in root.iter(qn('w:drawing')):
                for blip in drawing.iter(qn('a:blip')):
                    old_rid=blip.get(qn('r:embed'))
                    if not old_rid or old_rid not in part.related_parts:continue
                    old=part.related_parts[old_rid]
                    if not hasattr(old,'image') or old.image.px_width<=1 or old.image.px_height<=1:continue
                    ext=next(drawing.iter(qn('wp:extent')),None)
                    if ext is None:raise ValueError('El espacio del logo no tiene dimensiones editables.')
                    width,height=fit(image,int(ext.get('cx')),int(ext.get('cy')))
                    rid,_=part.get_or_add_image(str(path));blip.set(qn('r:embed'),rid)
                    ext.set('cx',str(width));ext.set('cy',str(height))
                    for xfrm in drawing.iter(qn('a:xfrm')):
                        for shape_ext in xfrm.findall(qn('a:ext')):
                            shape_ext.set('cx',str(width));shape_ext.set('cy',str(height))
                    for crop in list(drawing.iter(qn('a:srcRect'))):crop.getparent().remove(crop)
                    count+=1
    if address and not addresses:raise ValueError('No se encontró el campo de dirección en el encabezado.')
    if logo and not count:raise ValueError('No se encontró el logo del encabezado.')
    return count


def insert_signature(paragraph,path,width=Cm(2.2),height=Cm(.7)):
    path,image=validated_image(path,'la firma')
    width,height=fit(image,int(width),int(height))
    paragraph.add_run().add_picture(str(path),width=width,height=height)


def apply_signatures(doc,pid,people,representative):
    """Never create table columns or place signatures into phone/email cells."""
    inserted=0
    for table in doc.tables:
        head=[c.text.strip().upper() for c in cells(table.rows[0])]
        if 'FIRMA' not in head:continue
        idx=head.index('FIRMA')
        rows=table.rows[1:-1] if head[0]=='TOTAL DE COPIAS' else table.rows[1:]
        for row,person in zip(rows,people):
            if person.get('firma'):
                c=cells(row)[idx]
                width=min(int(Cm(2.2)),max(int(Cm(.3)),int(c.width or Cm(2.5))-int(Cm(.3))))
                insert_signature(c.paragraphs[0],person['firma'],width=width);inserted+=1
    rep=people[representative]
    if rep.get('firma'):
        if pid=='P01':
            target=next((p for p in doc.tables[2].cell(0,0).paragraphs if p.text.strip().startswith('Firma:')),None)
            if target is None:raise ValueError('No se encontró el espacio de firma de quien entrega.')
            cell_like_paragraph(target,'Firma: ');insert_signature(target,rep['firma'],width=Cm(3));inserted+=1
        elif pid=='P03':
            paragraphs=doc.paragraphs
            index=next((i for i,p in enumerate(paragraphs) if p.text==rep['nombre']),None)
            if index is None or index==0 or paragraphs[index-1].text.strip():
                raise ValueError('No se encontró el espacio previo al nombre para la firma de la nota.')
            insert_signature(paragraphs[index-1],rep['firma'],width=Cm(3),height=Cm(1));inserted+=1
    return inserted
