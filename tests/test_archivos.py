import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from docx import Document
from openpyxl import Workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from motor import Base
from sencillo import catalog

class FilesTests(unittest.TestCase):
    def test_missing_catalog_explains_extraction(self):
        with tempfile.TemporaryDirectory() as td,patch('sencillo.ROOT',Path(td)):
            with self.assertRaisesRegex(ValueError,'Extrae todo'):catalog()
    def test_wrong_sheet_explains_catalog_requirement(self):
        with tempfile.TemporaryDirectory() as td,patch('sencillo.ROOT',Path(td)):
            w=Workbook();w.save(Path(td)/'Base_SASISOPA.xlsx')
            with self.assertRaisesRegex(ValueError,'hoja Plantillas'):catalog()
    def test_blank_filename_rejected_with_readable_message(self):
        with tempfile.TemporaryDirectory() as td,patch('sencillo.ROOT',Path(td)):
            w=Workbook();s=w.active;s.title='Plantillas';s.append(['plantilla_id','archivo_word'])
            for i in range(1,8):s.append([f'P{i:02}',None if i==1 else f'p{i}.docx'])
            w.save(Path(td)/'Base_SASISOPA.xlsx')
            with self.assertRaisesRegex(ValueError,'archivo vacío'):catalog()
    def test_advanced_write_failure_preserves_existing_user_file(self):
        base=Base.__new__(Base)
        base.data={'Generar':[{'generar':'Sí','nombre_salida':'uno'},{'generar':'Sí','nombre_salida':'dos'}]}
        base.render=lambda job:Document()
        from docx.document import Document as DocClass
        original=DocClass.save;count=0
        def fail_second(doc,path):
            nonlocal count
            count+=1
            if count==2:raise OSError('Fallo simulado')
            return original(doc,path)
        with tempfile.TemporaryDirectory() as td:
            out=Path(td);keep=out/'conservar.docx';keep.write_bytes(b'Datos previos')
            with patch.object(DocClass,'save',fail_second),self.assertRaisesRegex(OSError,'simulado'):base.generate(out)
            self.assertEqual(keep.read_bytes(),b'Datos previos')
            self.assertEqual({p.name for p in out.iterdir()},{'conservar.docx'})
    def test_advanced_preexisting_document_not_changed(self):
        base=Base.__new__(Base);base.data={'Generar':[{'generar':'Sí','nombre_salida':'uno'}]};base.render=lambda job:Document()
        with tempfile.TemporaryDirectory() as td:
            out=Path(td);existing=out/'uno.docx';existing.write_bytes(b'No sobrescribir')
            with self.assertRaisesRegex(ValueError,'existe'):base.generate(out)
            self.assertEqual(existing.read_bytes(),b'No sobrescribir')

if __name__=='__main__':unittest.main()
