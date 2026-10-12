import os
import sys
import tempfile
import unittest
import zipfile
import hashlib
from pathlib import Path
from datetime import datetime
from types import SimpleNamespace
from unittest.mock import patch
from docx import Document
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from sencillo import basic_generate,catalog
from motor import export_pdf

class BasicTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.templates=Path(os.environ.get('SASISOPA_TEST_TEMPLATES','plantillas'))
        if len(list(cls.templates.glob('*.docx')))!=7:
            raise unittest.SkipTest('Se requieren las siete plantillas locales; no están publicadas.')
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.out=Path(self.tmp.name)/'lote'
        self.people=[{'nombre':f'PERSONA FICTICIA {i}','puesto':f'PUESTO {i}'} for i in range(1,7)]
    def generate(self):
        return basic_generate(self.templates,self.out,datetime(2027,2,4),'E-99999','ESTACIÓN FICTICIA',self.people,5,4)
    def test_seven_documents_structure_media_headers_and_fixed_content(self):
        before={p:hashlib.sha256(p.read_bytes()).hexdigest() for p in self.templates.glob('*.docx')}
        paths=self.generate();self.assertEqual(len(paths),7)
        for path in paths:
            with self.subTest(path=path.name):
                d=Document(path);source_name=dict(catalog())[path.name.split('_')[0]]
                o=Document(self.templates/source_name)
                self.assertEqual([len(t.rows) for t in d.tables],[len(t.rows) for t in o.tables])
                self.assertEqual(len(d.sections),len(o.sections))
                body='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
                self.assertIn('PERSONA FICTICIA',body)
                self.assertNotIn('13/ENE',body);self.assertNotIn('16/01/2025',body)
                historical=lambda doc:sum('30-07-2018' in c.text for t in doc.tables for r in t.rows for c in r.cells)
                self.assertEqual(historical(d),historical(o))
                # Meeting times and objectives retain their original text.
                for t in o.tables:
                    if len(t.rows)>=3 and t.rows[0].cells[0].text=='Tipo de Reunión:':
                        self.assertIn(t.rows[2].cells[1].text,body)
                with zipfile.ZipFile(path) as z,zipfile.ZipFile(self.templates/source_name) as original:
                    for name in original.namelist():
                        if name.startswith('word/media/'):self.assertEqual(z.read(name),original.read(name))
                    header=b''.join(z.read(n) for n in z.namelist() if n.startswith('word/header') and n.endswith('.xml'))
                    self.assertIn(b'E-99999',header);self.assertNotIn(b'E-02585',header)
        self.assertTrue(all(hashlib.sha256(p.read_bytes()).hexdigest()==h for p,h in before.items()))
    def test_overwrite_blocked(self):
        self.generate()
        with self.assertRaisesRegex(ValueError,'ya existe'):self.generate()
    def test_capacity_leaves_no_outputs(self):
        self.people=[{'nombre':f'PERSONA {i}','puesto':'PUESTO'} for i in range(12)]
        with self.assertRaisesRegex(ValueError,'admite'):self.generate()
        self.assertFalse(self.out.exists())
    def test_duplicate_people_blocked(self):
        self.people[1]=self.people[0].copy()
        with self.assertRaisesRegex(ValueError,'duplicados'):self.generate()
        self.assertFalse(self.out.exists())

class PDFLifecycleTests(unittest.TestCase):
    def test_com_initialization_cleanup_and_missing_pdf_detection(self):
        events=[]
        document=SimpleNamespace(ExportAsFixedFormat=lambda *a:events.append('export'),Close=lambda *a:events.append('close'))
        word=SimpleNamespace(Documents=SimpleNamespace(Open=lambda *a,**kw:document),Quit=lambda:events.append('quit'))
        modules={'pythoncom':SimpleNamespace(CoInitialize=lambda:events.append('init'),CoUninitialize=lambda:events.append('uninit')),
                 'win32com':SimpleNamespace(client=SimpleNamespace(DispatchEx=lambda *a:word)),
                 'win32com.client':SimpleNamespace(DispatchEx=lambda *a:word)}
        with tempfile.TemporaryDirectory() as td,patch.dict(sys.modules,modules):
            path=Path(td)/'prueba.docx';path.write_bytes(b'test')
            with self.assertRaisesRegex(RuntimeError,'no creó'):export_pdf([path])
        self.assertEqual(events,['init','export','close','quit','uninit'])
    def test_com_initialization_failure_still_releases_thread(self):
        events=[]
        def fail(*args):raise RuntimeError('Word no disponible')
        client=SimpleNamespace(DispatchEx=fail)
        modules={'pythoncom':SimpleNamespace(CoInitialize=lambda:events.append('init'),CoUninitialize=lambda:events.append('uninit')),'win32com':SimpleNamespace(client=client),'win32com.client':client}
        with patch.dict(sys.modules,modules):
            with self.assertRaisesRegex(RuntimeError,'Word no disponible'):export_pdf([])
        self.assertEqual(events,['init','uninit'])

if __name__=='__main__':unittest.main()
