import os
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from datetime import datetime
from PIL import Image
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from sencillo import basic_generate,catalog
from firma_canvas import signature_png

class CaptureTests(unittest.TestCase):
    def test_blank_signature_rejected(self):
        with self.assertRaisesRegex(ValueError,'Dibuja'):signature_png([])
    def test_drawn_signature_png_has_ink_and_transparency(self):
        image=signature_png([[(20,30),(60,80),(180,45)]])
        self.assertEqual(image.size,(600,200));self.assertEqual(image.mode,'RGBA')
        self.assertIsNotNone(image.getchannel('A').getbbox())
        self.assertEqual(image.getpixel((599,199))[3],0)
    def test_drawing_outside_canvas_rejected(self):
        with self.assertRaisesRegex(ValueError,'fuera'):signature_png([[(0,0),(800,20)]])

class BrandingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from catalogo_activo import catalog_state
        if catalog_state()['status']!='ready':raise unittest.SkipTest('El catálogo anterior de Word fue retirado; la nueva base está pendiente de adaptar.')
        cls.templates=Path(os.environ.get('SASISOPA_TEST_TEMPLATES','plantillas'))
        if len(list(cls.templates.glob('*.docx')))!=7:raise unittest.SkipTest('Faltan plantillas locales.')
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.logo=self.root/'logo.png';Image.new('RGB',(200,200),(0,100,190)).save(self.logo)
        self.signature=self.root/'firma.png';signature_png([[(20,40),(100,100),(300,50)]]).save(self.signature)
        self.people=[{'nombre':f'PERSONA FICTICIA {i}','puesto':'PUESTO','firma':str(self.signature)} for i in range(6)]
    def run_batch(self,**kwargs):
        return basic_generate(self.templates,self.root/'lote',datetime(2028,2,29),'E-99999','ESTACIÓN FICTICIA',self.people,5,4,**kwargs)
    def test_logo_address_revision_and_existing_signature_slots(self):
        paths=self.run_batch(address='DIRECCIÓN FICTICIA 100',logo=str(self.logo),revision='3',revision_date=datetime(2027,5,12),signatures=True)
        for path in paths:
            pid=path.name.split('_')[0];d=Document(path);o=Document(self.templates/dict(catalog())[pid])
            self.assertEqual([len(t.rows) for t in d.tables],[len(t.rows) for t in o.tables])
            found_logo=False
            for rel in d.part.rels.values():
                if rel.reltype.endswith('/header'):
                    root=rel.target_part.element
                    header=''.join(t.text or '' for t in root.iter(qn('w:t')))
                    if not header.strip():continue  # Original first/even-page headers may be intentionally empty.
                    self.assertIn('DIRECCIÓN FICTICIA 100',header)
                    if pid!='P03':
                        self.assertIn('MAYO2027',header)
                        self.assertIn('.R3',header)
                        self.assertNotIn('.R0',header)
                    for drawing in root.iter(qn('w:drawing')):
                        for blip in drawing.iter(qn('a:blip')):
                            part=rel.target_part.related_parts[blip.get(qn('r:embed'))]
                            if part.blob==self.logo.read_bytes():
                                extent=next(drawing.iter(qn('wp:extent')))
                                self.assertEqual(extent.get('cx'),extent.get('cy'))
                                found_logo=True
            self.assertTrue(found_logo,pid)
            # Only the explicit signature spaces in these originals receive images.
            signature_blips=[]
            for blip in d.element.iter(qn('a:blip')):
                part=d.part.related_parts[blip.get(qn('r:embed'))]
                if part.blob==self.signature.read_bytes():signature_blips.append(blip)
            expected={'P01':7,'P03':1,'P04':6}.get(pid,0)
            self.assertEqual(len(signature_blips),expected,pid)
            self.assertEqual(len(d.paragraphs),len(o.paragraphs))
    def test_missing_image_rejected_before_outputs(self):
        with self.assertRaisesRegex(ValueError,'No se encuentra'):self.run_batch(logo=str(self.root/'falta.png'))
        self.assertFalse((self.root/'lote').exists())
    def test_invalid_image_rejected_before_outputs(self):
        bad=self.root/'falsa.png';bad.write_text('not an image')
        with self.assertRaisesRegex(ValueError,'No se puede leer'):self.run_batch(logo=str(bad))
        self.assertFalse((self.root/'lote').exists())
    def test_corrupt_png_pixel_data_rejected_before_outputs(self):
        damaged=self.root/'dañada.png'
        data=bytearray(self.logo.read_bytes())
        data[data.index(b'IDAT')+5]^=1
        damaged.write_bytes(data)
        with self.assertRaisesRegex(ValueError,'No se puede leer'):self.run_batch(logo=str(damaged))
        self.assertFalse((self.root/'lote').exists())
    def test_invalid_revision_rejected_without_outputs(self):
        with self.assertRaisesRegex(ValueError,'número entero'):self.run_batch(revision='3.5')
        self.assertFalse((self.root/'lote').exists())
    def test_no_signature_selection_keeps_original_blank_fields(self):
        paths=self.run_batch(signatures=False)
        for path in paths:
            d=Document(path)
            for blip in d.element.iter(qn('a:blip')):
                self.assertNotEqual(d.part.related_parts[blip.get(qn('r:embed'))].blob,self.signature.read_bytes())

if __name__=='__main__':unittest.main()
