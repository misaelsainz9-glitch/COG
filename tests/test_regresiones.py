import sys
import os
import tempfile
import shutil
import unittest
from pathlib import Path
from datetime import datetime
from unittest.mock import patch
from docx import Document
from docx.oxml.ns import qn
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from motor import replace,cell,fill_session_people,positive_integer
from sencillo import basic_generate,catalog

class XMLRegressionTests(unittest.TestCase):
    def test_replacement_across_runs_and_multiple_matches(self):
        d=Document();p=d.add_paragraph()
        p.add_run('AA E-02').bold=True;p.add_run('585 BB E-02585 CC')
        replace(d,'E-02585','E-02585-NUEVA')
        self.assertEqual(p.text,'AA E-02585-NUEVA BB E-02585-NUEVA CC')
        self.assertTrue(p.runs[0].bold)
    def test_replacement_does_not_create_headers(self):
        d=Document();d.add_paragraph('ESTACION')
        before=[r.reltype for r in d.part.rels.values()]
        replace(d,'ESTACION','NUEVA')
        self.assertEqual([r.reltype for r in d.part.rels.values()],before)
    def test_cell_uses_existing_paragraphs_without_duplicate_breaks(self):
        d=Document();c=d.add_table(rows=1,cols=1).cell(0,0)
        c.paragraphs[0].add_run('Original').bold=True;c.add_paragraph('Otra');c.add_paragraph('Más')
        cell(c,'Uno\nDos\nTres')
        self.assertEqual([p.text for p in c.paragraphs],['Uno','Dos','Tres'])
        self.assertTrue(c.paragraphs[0].runs[0].bold)
    def test_quantities_never_truncate_decimals(self):
        self.assertEqual(positive_integer(2.0,'Copias'),2)
        for bad in (1.5,0,-1,'NaN','Infinity','abc',None):
            with self.subTest(value=bad),self.assertRaises(ValueError):positive_integer(bad,'Copias')
    def test_cell_preserves_inline_drawing_reference(self):
        import io,base64
        image=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jR1sAAAAASUVORK5CYII=')
        d=Document();c=d.add_table(rows=1,cols=1).cell(0,0)
        c.paragraphs[0].add_run().add_picture(io.BytesIO(image))
        cell(c,'Nuevo texto')
        self.assertEqual(len(c._tc.findall('.//'+qn('w:drawing'))),1)
        self.assertEqual(c.text,'Nuevo texto')
    def test_modified_list_rejected(self):
        d=Document();c=d.add_table(rows=1,cols=1).cell(0,0)
        c.text='Un formato diferente'
        with self.assertRaisesRegex(ValueError,'modificado'):fill_session_people(c,'E-99999',['Nombre'])

class TemplateRegressionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.source=Path(os.environ.get('SASISOPA_TEST_TEMPLATES','plantillas'))
        if len(list(cls.source.glob('*.docx')))!=7:raise unittest.SkipTest('Faltan plantillas locales.')
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.addCleanup(self.tmp.cleanup)
        self.root=Path(self.tmp.name)
        self.people=[{'nombre':f'PERSONA {i}','puesto':'PUESTO'} for i in range(6)]
    def run_batch(self,source=None):
        return basic_generate(source or self.source,self.root/'lote',datetime(2028,2,29),'E-99999','RAZÓN SOCIAL',self.people,5,4)
    def test_session_preserves_all_original_paragraph_slots(self):
        name=dict(catalog())['P02'];original=Document(self.source/name)
        paths=self.run_batch();p=next(p for p in paths if p.name.startswith('P02_'));doc=Document(p)
        old=original.tables[1].cell(3,1).paragraphs;new=doc.tables[1].cell(3,1).paragraphs
        self.assertEqual(len(new),len(old))
        self.assertEqual(sum(len(p._p.findall('.//'+qn('w:br'))) for p in new),sum(len(p._p.findall('.//'+qn('w:br'))) for p in old))
        self.assertEqual(new[2].text,'1. PERSONA 0');self.assertEqual(new[12].text,'6. PERSONA 5')
        self.assertEqual(new[-1].text,old[-1].text)
        self.assertTrue(all('20280229' in p.name for p in paths))
    def test_modified_template_rejected_without_outputs(self):
        source=self.root/'plantillas';shutil.copytree(self.source,source)
        file=source/dict(catalog())['P02'];d=Document(file);d.add_paragraph('Cambio');d.save(file)
        with self.assertRaisesRegex(ValueError,'modificada'):self.run_batch(source)
        self.assertFalse((self.root/'lote').exists())
    def test_failed_save_removes_only_new_batch(self):
        keep=self.root/'conservar.txt';keep.write_text('No tocar')
        from docx.document import Document as DocClass
        original=DocClass.save;count=0
        def save(doc,path):
            nonlocal count
            count+=1
            if count==2:raise OSError('Fallo de escritura simulado')
            return original(doc,path)
        with patch.object(DocClass,'save',save):
            with self.assertRaisesRegex(OSError,'simulado'):self.run_batch()
        self.assertFalse((self.root/'lote').exists());self.assertEqual(keep.read_text(),'No tocar')

if __name__=='__main__':unittest.main()
